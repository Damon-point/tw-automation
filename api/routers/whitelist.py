"""
WhiteList · DQA — 8080 의 /whitelist(메일함 · 신청 · 등록 내역 · 규칙 · 고객 포털)와
처리 팝업 세 가지(일반 등록 · AI 해석(대량/첨부) · DQA FOTA Test)를 옮긴 API.

업무 로직은 backend/app/whitelist/* · app/pipeline/dqa_* 를 그대로 부른다. 화면 쪽 상태
(팝업이 들고 있던 파싱값·해석 결과·포털 조회)는 이 프로세스 메모리에 스레드별로 둔다 —
화면을 닫았다 다시 열어도 이어서 볼 수 있게.

외부 쓰기(FOTA 엑셀·FOTA DB·FOTA 웹·Dropbox·메일)는 app.shadow 가 검증 모드에서 막는다.
그래서 8080 과 같은 인자(allow_write=True 등)로 부른다.
"""
from __future__ import annotations

import os
import re
import threading
import time
from datetime import datetime
from io import BytesIO
from pathlib import Path

from fastapi import APIRouter, File, HTTPException, Request, UploadFile
from fastapi.responses import FileResponse, Response
from pydantic import BaseModel

from common import admin, user

router = APIRouter(prefix="/api/wl", tags=["whitelist"])

MODELS = ["PM451", "PM45", "PM560", "PM67", "PM85", "PM86", "PM90", "PM95"]
OS_LEVELS = [6, 7, 8, 9, 10, 11, 12, 13, 14, 15, 16]
STATUS_KO = {"pending": "대기 중", "confirmed": "승인됨", "rejected": "거절됨",
             "deleted": "삭제됨", "handled": "수동 처리(등록 아님)"}
_CHUNK = 200


def _S():
    from app.whitelist import service as S

    return S


def _bad(msg: str, code: int = 400):
    raise HTTPException(code, msg)


# ── 메일함: 스레드 목록 ────────────────────────────────────────────────────────
# 목록 계산은 메일 전량 분류 + FOTA 조회라 첫 회 수 초가 걸린다. 처리 완료·작업 아님까지
# 전부 포함한 목록을 한 번 계산해 두고, 화면의 체크박스는 거기서 거른다.

_cache = {"at": 0.0, "groups": None}
_cache_lock = threading.Lock()
_TTL = 20.0


def _compute() -> list[dict]:
    from app.pages.whitelist_mail import _threads

    try:
        groups = _threads({"q": "", "show_done": True, "show_notwl": True})
    finally:
        with _cache_lock:
            _cache["refreshing"] = False
    with _cache_lock:
        _cache.update(at=time.time(), groups=groups)
    return groups


def _all_groups(force: bool = False) -> list[dict]:
    """목록. 오래됐으면 지금 것을 돌려주고 뒤에서 다시 계산한다(2~3초) — 처리 직후만 기다린다."""
    with _cache_lock:
        cached = _cache["groups"]
        stale = time.time() - _cache["at"] >= _TTL
        if cached is not None and not force:
            if stale and not _cache.get("refreshing"):
                _cache["refreshing"] = True
                threading.Thread(target=_compute, daemon=True).start()
            return cached
    return _compute()


def _invalidate() -> None:
    with _cache_lock:
        _cache["groups"] = None


def _tid(g: dict) -> str:
    return str(g["key"]).rsplit("#", 1)[-1]


def _find(tid: str) -> dict:
    for g in _all_groups():
        if _tid(g) == str(tid):
            return g
    for g in _all_groups(force=True):
        if _tid(g) == str(tid):
            return g
    raise HTTPException(404, "메일 스레드를 찾을 수 없습니다")


def _when(received_at: str) -> str:
    from app.pages.whitelist_mail import _when as w

    return w(received_at)


def _kind(g: dict) -> str:
    """처리 방식 — dqa(DQA FOTA Test) · organic(대량/첨부, AI 해석) · simple(일반 등록) · none."""
    from app.pages.whitelist_bulk import BULK_THRESHOLD
    from app.pipeline import dqa_fota as D

    rep = g["rep"]
    if D.is_fota_test_mail(rep.get("sender", ""), rep.get("subject", ""), rep.get("body_text", "")):
        return "dqa"
    att = rep.get("attachment_names") or ""
    if len(g["serials"]) > BULK_THRESHOLD or "[첨부메일]" in att \
            or any(x in att.lower() for x in (".xlsx", ".xls", ".csv")):
        return "organic"
    return "simple" if g["serials"] else "none"


def _state(g: dict) -> str:
    if g.get("notwl"):
        return "notwl"
    if g.get("done"):
        return "done"
    if g.get("reopened"):
        return "reopened"
    return "open"


def _summary(g: dict) -> dict:
    rep = g["rep"]
    v = g.get("verdict") or {}
    return {
        "tid": _tid(g), "subject": rep.get("subject", ""), "sender": rep.get("sender", ""),
        "when": _when(rep.get("received_at", "")), "count": g.get("count", 1),
        "serial_count": len(g["serials"]), "dup_count": len(g.get("dups") or {}),
        "attachments": [a for a in re.split(r"[\n,]", rep.get("attachment_names") or "") if a.strip()][:6],
        "judged": bool(v), "model": v.get("model") or "", "from_os": str(v.get("from_os") or ""),
        "to_os": str(v.get("to_os") or ""), "reason": v.get("reason") or "",
        "state": _state(g), "processed": bool(g.get("processed")), "kind": _kind(g),
        "snippet": " ".join((rep.get("body_text") or "").split())[:160],
    }


@router.get("/threads")
def threads(request: Request, show_done: bool = False, show_notwl: bool = False, q: str = ""):
    user(request)
    groups = _all_groups()
    open_n = sum(1 for g in groups if not g.get("done") and not g.get("notwl"))
    out = [g for g in groups if (show_done or not g.get("done")) and (show_notwl or not g.get("notwl"))]
    ql = q.strip().lower()
    if ql:
        out = [g for g in out if ql in g["rep"]["subject"].lower() or ql in g["rep"]["sender"].lower()
               or any(ql in s.lower() for s in g["serials"])]
    return {"open": open_n, "threads": [_summary(g) for g in out]}


@router.get("/threads/{tid}")
def thread_detail(tid: str, request: Request):
    u = user(request)
    from app.pages.whitelist_mail import _sender_addr_name
    from app.whitelist import mailer as WM

    S = _S()
    g = _find(tid)
    rep = g["rep"]
    addr, name = _sender_addr_name(rep.get("sender", ""))
    members = sorted(g.get("members") or [rep], key=lambda r: str(r.get("id")))
    from app.whitelist import mail_triage as T

    members = sorted(members, key=T.received_key, reverse=True)
    return {
        **_summary(g),
        "is_admin": u.get("role") == "admin",
        "addr": addr, "name": name,
        "lang": WM.detect_lang(f"{rep.get('subject', '')}\n{rep.get('body_text', '')}"),
        "auto_category": S.category_for_sender(rep.get("sender", "")),
        "serials": [{"sn": s, "dup": (g.get("dups") or {}).get(s)} for s in g["serials"]],
        "members": [{"id": m.get("id"), "subject": m.get("subject", ""), "sender": m.get("sender", ""),
                     "when": _when(m.get("received_at", "")),
                     "body": T.strip_quotes(m.get("body_text") or "")[:6000],
                     "attachments": m.get("attachment_names") or ""} for m in members],
    }


@router.post("/sync")
def sync(request: Request):
    u = user(request)
    from app import background as B
    from app.pages.whitelist_mail import auto_triage_new

    try:
        n = B.locked_sync(u["id"])
    except Exception as e:  # noqa: BLE001
        _bad(f"동기화 오류: {e}")
    try:
        res = auto_triage_new()
    except Exception as e:  # noqa: BLE001
        res = {"error": str(e)}
    _invalidate()
    return {"added": n, "triage": {k: res.get(k) for k in ("asked", "pending", "blocked", "error")}}


@router.post("/triage")
def triage(request: Request):
    user(request)
    from app.pages.whitelist_mail import auto_triage_new

    res = auto_triage_new()
    _invalidate()
    return {k: res.get(k) for k in ("asked", "pending", "blocked", "error", "est_cost")}


class FlagIn(BaseModel):
    on: bool


@router.post("/threads/{tid}/done")
def set_done(tid: str, body: FlagIn, request: Request):
    u = user(request)
    from app.pages.whitelist_mail import _sender_addr_name
    from app.whitelist import mail_state as MS
    from app.whitelist import mail_triage as T

    S = _S()
    g = _find(tid)
    MS.set_status(g["key"], MS.DONE if body.on else MS.OPEN, by=u["username"],
                  note="수동 처리 완료" if body.on else "", rep_ts=T.newest_ts(g))
    try:
        rep = g["rep"]
        if body.on:
            a, n = _sender_addr_name(rep.get("sender", ""))
            S.record_manual_done(mail={"mail_id": rep["id"], "requester_addr": a, "requester_name": n,
                                       "message_id": rep.get("message_id", "")},
                                 serials=list(g.get("serials") or []), processed_by=u["username"],
                                 verdict=g.get("verdict") or {})
        else:
            S.remove_manual_done(rep["id"])
    except Exception:  # noqa: BLE001
        pass
    _invalidate()
    return {"ok": True, "message": "처리 완료로 표시했습니다" if body.on else "다시 열었습니다"}


@router.post("/threads/{tid}/notwl")
def set_notwl(tid: str, body: FlagIn, request: Request):
    u = user(request)
    from app.whitelist import mail_state as MS
    from app.whitelist import mail_triage as T

    g = _find(tid)
    MS.set_status(g["key"], MS.NOTWL if body.on else MS.OPEN, by=u["username"],
                  note="WhiteList 작업 아님" if body.on else "", rep_ts=T.newest_ts(g))
    _invalidate()
    return {"ok": True, "message": "작업 아님으로 내렸습니다" if body.on else "목록에 다시 올렸습니다"}


def _close_thread(g: dict, me: str, note: str) -> None:
    from app.whitelist import mail_state as MS
    from app.whitelist import mail_triage as T

    try:
        MS.set_status(g["key"], MS.DONE, by=me, note=note, rep_ts=T.newest_ts(g))
    except Exception:  # noqa: BLE001
        pass
    _invalidate()


class SnsIn(BaseModel):
    sns: list[str]


@router.post("/fota-status")
def fota_status(body: SnsIn, request: Request):
    """각 SN 의 FOTA 화이트리스트 등록 여부(읽기)."""
    user(request)
    from app.pages.whitelist_mail import _fota_status

    return {"status": _fota_status([s for s in body.sns if s][:400])}


# ── 일반 등록(메일 1건) ─────────────────────────────────────────────────────────

class DevSpecIn(BaseModel):
    sns: list[str]
    model: str = ""
    cur: str = ""
    upg: str = ""
    category: str = ""
    ostype_override: str = ""
    registered: list[str] = []


def _upload_rows(devs: list[dict], registered: set[str] | None = None) -> list[dict]:
    from app.pages.whitelist_bulk import upload_rows

    return upload_rows(devs, {s: True for s in (registered or set())})


def _mail_devices(spec: DevSpecIn) -> list[dict]:
    S = _S()
    cat = spec.category or S.DEFAULT_CATEGORY
    ov = (spec.ostype_override or "").strip()
    out = []
    for s in spec.sns:
        r = S.resolve_os(spec.model, spec.cur, spec.upg, cat)
        c = str(r["from"]) if r else spec.cur
        u = str(r["to"]) if r else spec.upg
        out.append(S.build_device(s, c, u, spec.model, cat, ostype_override=ov))
    return out


@router.post("/threads/{tid}/devices")
def mail_devices(tid: str, body: DevSpecIn, request: Request):
    user(request)
    devs = _mail_devices(body)
    return {"rows": _upload_rows(devs, set(body.registered)),
            "valid": sum(1 for d in devs if d["valid"]), "total": len(devs)}


def _mail_ctx(g: dict, uid: int) -> dict:
    from app.pages.whitelist_mail import _sender_addr_name
    from app.whitelist import mailer as WM

    rep = g["rep"]
    a, n = _sender_addr_name(rep.get("sender", ""))
    return {"mail_id": rep["id"], "user_id": uid, "requester_addr": a, "requester_name": n,
            "lang": rep.get("lang") or WM.detect_lang(f"{rep.get('subject', '')}\n{rep.get('body_text', '')}"),
            "message_id": rep.get("message_id", "")}


class RegisterIn(DevSpecIn):
    charge_type: str = "무상"
    free_reason: str = ""
    memo: str = ""
    write: bool = True


@router.post("/threads/{tid}/register")
def mail_register(tid: str, body: RegisterIn, request: Request):
    u = admin(request)
    S = _S()
    g = _find(tid)
    goods = [d for d in _mail_devices(body) if d["valid"]]
    if not goods:
        _bad("등록 가능한 단말이 없습니다")
    reg = set(body.registered)
    fresh = [d["sn"] for d in goods if d["sn"] not in reg]
    dup = S.existing_sns(fresh, fresh=True) if fresh else {}
    if dup:
        _bad("이미 신청된 단말: " + ", ".join(dup), 409)
    paid = body.charge_type == "유상"
    res = S.register_from_mail(
        devices=goods, processed_by=u["username"], charge_type=body.charge_type,
        category=body.category or S.DEFAULT_CATEGORY,
        free_reason="" if paid else body.free_reason.strip(), memo=body.memo.strip(),
        mail=_mail_ctx(g, u["id"]), allow_duplicates=True, allow_write=body.write)
    if not res.get("ok"):
        _bad(res.get("msg", "등록 실패"))
    f = res.get("fota") or {}
    if res.get("dry"):
        return {"dry": True, "fota": _fota_brief(f), "request": None, "verify_items": []}
    _invalidate()
    return {"dry": False, "fota": _fota_brief(f), "request": res["request"],
            "verify_items": [{"sn": d["sn"], "endDate": d.get("endDate", ""), "ostype": d.get("ostype", ""),
                              "api": str(d.get("api") or "")} for d in goods],
            "devices": [{"sn": d["sn"]} for d in goods]}


def _fota_brief(f: dict) -> dict:
    return {"ok": f.get("ok"), "total": f.get("total", 0), "success": f.get("success", 0),
            "msg": f.get("msg", ""), "dry": f.get("dryRun"),
            "fails": [d for d in f.get("details", []) if not d.get("ok")][:30]}


class DateIn(BaseModel):
    sns: list[str]
    date: str
    write: bool = True


@router.post("/threads/{tid}/change-date")
def mail_change_date(tid: str, body: DateIn, request: Request):
    u = admin(request)
    S = _S()
    g = _find(tid)
    nd = body.date.strip()
    if not re.match(r"^\d{4}-\d{2}-\d{2}$", nd):
        _bad("날짜는 YYYY-MM-DD 형식입니다")
    sns = [s for s in body.sns if s]
    res = S.change_dates_db(sns=sns, new_date=nd, processed_by=u["username"],
                            mail=_mail_ctx(g, u["id"]), dry_run=not body.write)
    out = {"ok": bool(res.get("ok")), "dry": bool(res.get("dry")), "lines": res.get("lines", []),
           "request": res.get("request"), "missing": res.get("missing", []), "verify_items": []}
    if res.get("ok") and not res.get("dry"):
        y, m, d = nd.split("-")
        out["verify_items"] = [{"sn": s, "endDate": f"{int(y) + 1}-{m}-{d}"} for s in sns]
        _invalidate()
    return out


@router.post("/requests/{req_id}/po")
async def upload_po(req_id: str, request: Request, file: UploadFile = File(...)):
    user(request)
    S = _S()
    data = await file.read()
    p = S.save_po_file(req_id, file.filename or "po", data)
    S.attach_po(req_id, p)
    return {"ok": True, "path": os.path.basename(p)}


class VerifyIn(BaseModel):
    items: list[dict]


@router.post("/verify")
def verify(body: VerifyIn, request: Request):
    u = user(request)
    from app.whitelist import fota as WF

    res = WF.verify_registration(u["id"], body.items)
    rows = res.get("rows") or []
    bad = [r for r in rows if not r.get("ok")]
    shown = rows if len(rows) <= 300 else (bad or rows)[:300]
    return {**{k: res.get(k) for k in ("ok", "total", "matched", "missing", "mismatched", "error")},
            "rows": shown, "hidden": len(rows) - len(shown)}


# ── 회신 ───────────────────────────────────────────────────────────────────────

class DraftIn(BaseModel):
    kind: str = ""            # '' | '등록' | '등록일 변경' | 'dqa'
    new_date: str = ""
    sns: list[str] = []


def _reply_ctx(g: dict, with_cc: bool) -> dict:
    from app.user_mailer import reply_context

    rep = g["rep"]
    return (reply_context(rep.get("id"), with_cc=with_cc) or {}) if rep.get("id") else {}


@router.post("/threads/{tid}/reply-draft")
def reply_draft(tid: str, body: DraftIn, request: Request):
    user(request)
    from app.pages.whitelist_mail import _sender_addr_name
    from app.user_mailer import display_name_for
    from app.whitelist import mailer as WM

    g = _find(tid)
    rep = g["rep"]
    addr, name = _sender_addr_name(rep.get("sender", ""))
    if body.kind == "dqa":
        if not name and addr:
            name = display_name_for(addr)
        d = WM.build_dqa_reply_draft(to_name=name, subject_orig=rep.get("subject", ""))
        cc = _reply_ctx(g, True).get("cc", "")
    else:
        lang = WM.detect_lang(f"{rep.get('subject', '')}\n{rep.get('body_text', '')}")
        d = WM.build_reply_draft(to_name=name, devices=[{"sn": s} for s in (body.sns or g["serials"])],
                                 lang=lang, kind=body.kind, new_date=body.new_date)
        cc = ""
    return {"to": addr, "cc": cc if isinstance(cc, str) else ", ".join(cc or []), **d}


class SendIn(BaseModel):
    subject: str
    body: str
    with_cc: bool = False
    request_id: str = ""
    close: bool = True
    note: str = "회신 완료"


@router.post("/threads/{tid}/reply-preview")
def reply_preview(tid: str, body: SendIn, request: Request):
    u = user(request)
    from app.user_mailer import build_html_mail, get_send_identity

    g = _find(tid)
    ident = get_send_identity(u["id"])
    ctx = _reply_ctx(g, body.with_cc)
    doc = build_html_mail(body.body, quote=ctx.get("quote_html", ""),
                          signature=(ident.signature if ident else ""))
    return {"from": ident.email if ident else "", "html": doc}


@router.post("/threads/{tid}/reply")
def reply_send(tid: str, body: SendIn, request: Request):
    u = admin(request)
    from app.pages.whitelist_mail import _sender_addr_name
    from app.whitelist import mailer as WM

    S = _S()
    g = _find(tid)
    rep = g["rep"]
    addr, _n = _sender_addr_name(rep.get("sender", ""))
    if not addr:
        _bad("회신 받을 주소가 없습니다")
    ok, msg = WM.send_reply(u["id"], addr, body.subject, body.body,
                            orig_mail_id=rep.get("id"), with_cc=body.with_cc)
    if body.request_id:
        S.set_reply_status(body.request_id, "sent" if ok else "failed")
    if not ok:
        _bad(f"회신 발송 실패: {msg}")
    if body.close:
        _close_thread(g, u["username"], body.note)
    return {"ok": True, "message": "회신을 보냈습니다"}


@router.post("/threads/{tid}/close")
def close_thread(tid: str, request: Request, note: str = "처리 완료"):
    u = user(request)
    _close_thread(_find(tid), u["username"], note)
    return {"ok": True}


# ── AI 해석 처리(대량 · 첨부 · 양식 없는 요청) ─────────────────────────────────────

_organic: dict[str, dict] = {}
_TYPE_KO = {"date": "등록일 변경 요청", "new": "신규 등록 요청", "mixed": "혼합(판단 필요)", "unknown": "판별 실패"}


def _dedup(sns) -> list[str]:
    return list(dict.fromkeys(s.strip().upper() for s in sns if s and s.strip()))


def _org_load(g: dict, reinterpret: bool = False, instruction: str = "") -> dict:
    import email as _email

    from app.mail import attachments as ATT
    from app.mail import service as MS
    from app.whitelist import interpret as INT

    rep = g["rep"]
    att_dir = rep.get("attachment_dir", "")
    raw = ATT.load_raw(att_dir)
    src = "저장 원문"
    if raw is None:
        try:
            raw = MS.fetch_raw_by_uidl(rep.get("uidl", ""))
        except Exception:  # noqa: BLE001
            raw = None
        src = "서버 재조회"
    ctx: dict = {}
    if raw:
        msg = _email.message_from_bytes(raw)
        blocks = ATT.extract_blocks(msg)
        images = ATT.collect_body_images(msg)
        cached = None if (reinterpret or instruction) else ATT.load_plan(att_dir)
        if isinstance(cached, dict) and cached.get("blocks") is not None:
            interp, note = cached, f"{src} · 저장된 AI 해석"
        else:
            interp = INT.interpret_blocks(blocks, rep.get("subject", ""), images, instruction)
            if interp.get("blocks") or interp.get("overall"):
                ATT.save_plan(att_dir, {"overall": interp.get("overall", {}), "blocks": interp.get("blocks", [])})
            note = (f"{src} · AI 해석" + (f" · 본문 이미지 {len(images)}장" if images else "")
                    + (" · 관리자 지시 반영" if instruction else ""))
        ctx["blocks"] = blocks
        ctx["plan_by_idx"] = {b["index"]: b for b in interp.get("blocks", [])}
        ctx["overall"] = dict(interp.get("overall", {}) or {})
        ctx["note"] = (interp.get("error", "") + " · " + note).strip(" ·")
    else:
        ctx["blocks"] = [{"source": "본문(원문 없음 — 저장 SN 사용)", "text": "",
                          "serials": _dedup(g.get("serials") or [])}]
        ctx["plan_by_idx"], ctx["overall"] = {}, {}
        ctx["note"] = "원문을 찾지 못했습니다 — 아래 지시로 유형을 알려주세요."
    ctx["lookup"] = None
    _organic[_tid(g)] = ctx
    return ctx


def _org_ctx(g: dict) -> dict:
    return _organic.get(_tid(g)) or _org_load(g)


def _org_all(ctx) -> list[str]:
    return _dedup([s for b in ctx["blocks"] for s in b["serials"]])


def _org_new(ctx) -> list[str]:
    out: list[str] = []
    for i, b in enumerate(ctx["blocks"]):
        if (ctx["plan_by_idx"].get(i, {}).get("type") or "skip") != "skip":
            out += b["serials"]
    return _dedup(out) or _org_all(ctx)


def _org_view(g: dict, ctx: dict) -> dict:
    ov = ctx["overall"]
    return {
        "type": ov.get("request_type", ""), "type_label": _TYPE_KO.get(ov.get("request_type", ""), "미판별"),
        "reason": ov.get("reason", ""), "date": ov.get("date", ""), "note": ctx.get("note", ""),
        "all_sns": _org_all(ctx), "new_sns": _org_new(ctx),
        "auto_category": _S().category_for_sender(g["rep"].get("sender", "")),
        "blocks": [{"index": i, "source": b.get("source", ""), "serials": len(b.get("serials") or []),
                    "plan": ctx["plan_by_idx"].get(i, {})} for i, b in enumerate(ctx["blocks"])],
    }


@router.get("/organic/{tid}")
def organic(tid: str, request: Request):
    admin(request)
    g = _find(tid)
    return _org_view(g, _org_ctx(g))


class ReinterpretIn(BaseModel):
    instruction: str = ""
    force_type: str = ""


@router.post("/organic/{tid}/interpret")
def organic_interpret(tid: str, body: ReinterpretIn, request: Request):
    admin(request)
    g = _find(tid)
    if body.force_type:
        ctx = _org_ctx(g)
        ctx["overall"]["request_type"] = body.force_type
    else:
        ctx = _org_load(g, reinterpret=True, instruction=body.instruction.strip())
    return _org_view(g, ctx)


class OrgNewIn(BaseModel):
    category: str = ""
    picks: dict[str, dict] = {}          # 모델 → {from, to}
    manual: dict[str, dict] = {}         # S/N → {model, from, to}
    write: bool = False
    refresh_lookup: bool = False


def _org_lookup(g: dict, ctx: dict, refresh: bool = False):
    from app.whitelist import portal_lookup as PLK

    if refresh or ctx.get("lookup") is None:
        ctx["lookup"] = PLK.lookup(_org_new(ctx))
    return ctx["lookup"]


def _org_build(g, ctx, body: OrgNewIn):
    S = _S()
    look = _org_lookup(g, ctx, body.refresh_lookup)
    sns = _org_new(ctx)
    cat = body.category or S.DEFAULT_CATEGORY
    groups = look.by_model(sns)
    model_devs: dict[str, list[dict]] = {}
    for m, lst in groups.items():
        p = body.picks.get(m) or {}
        rules = [r for r in S.rules_for(m) if r["category"] == cat] or S.rules_for(m)
        if not p and len(rules) == 1:
            p = {"from": str(rules[0]["from"]), "to": str(rules[0]["to"])}
        model_devs[m] = [S.build_device(sn, p.get("from", ""), p.get("to", ""), m, cat) for sn in lst]
    manual = {sn: a for sn, a in body.manual.items() if sn in look.failed}
    man_devs = [S.build_device(sn, a.get("from", ""), a.get("to", ""), a.get("model", ""), cat)
                for sn, a in manual.items()]
    return look, sns, cat, groups, model_devs, manual, man_devs


@router.post("/organic/{tid}/new-preview")
def organic_new_preview(tid: str, body: OrgNewIn, request: Request):
    admin(request)
    S = _S()
    g = _find(tid)
    ctx = _org_ctx(g)
    look, sns, cat, groups, model_devs, manual, man_devs = _org_build(g, ctx, body)
    models = []
    for m, devs in model_devs.items():
        rules = [r for r in S.rules_for(m) if r["category"] == cat] or S.rules_for(m)
        rows = _upload_rows(devs)
        models.append({"model": m, "count": len(devs), "valid": sum(1 for d in devs if d["valid"]),
                       "rules": [{"from": r["from"], "to": r["to"], "ostype": r["ostype"]} for r in rules],
                       "rows": rows[:300]})
    from app.pages.whitelist_organic import unknown_rows

    unk = unknown_rows(look.failed, manual, cat)
    ready = sum(m["valid"] for m in models) + sum(1 for d in man_devs if d["valid"])
    return {"category": cat, "total": len(sns), "looked_up": len(sns) - len(look.failed),
            "error": look.error or "", "models": models, "unknown": unk[:600], "unknown_total": len(unk),
            "ready": ready, "rule_models": sorted({r["model"] for r in S.list_rules()})}


@router.post("/organic/{tid}/new-run")
def organic_new_run(tid: str, body: OrgNewIn, request: Request):
    u = admin(request)
    S = _S()
    g = _find(tid)
    ctx = _org_ctx(g)
    look, sns, cat, groups, model_devs, manual, man_devs = _org_build(g, ctx, body)
    devs, invalid, reasons = [], 0, []
    for d in [x for v in model_devs.values() for x in v] + man_devs:
        if d["valid"]:
            devs.append({"sn": d["sn"], "api": d["api"], "endDate": d["endDate"], "ostype": d["ostype"],
                         "partnumber": "", "model": d["model"], "cur": d["cur"], "upg": d["upg"]})
        else:
            invalid += 1
            rr = ", ".join(d["errors"])
            if rr not in reasons and len(reasons) < 3:
                reasons.append(rr)
    invalid += len([s for s in look.failed if s not in manual])
    parts = [f"{m} {(body.picks.get(m) or {}).get('from', '')}→{(body.picks.get(m) or {}).get('to', '')}"
             for m in groups]
    parts += sorted({f"{a['model']} {a['from']}→{a['to']}(수동)" for a in manual.values()})
    label = "신규 등록 · " + " · ".join(parts)
    if not devs:
        return {"dry": not body.write, "ok": False, "success": 0, "fails": [], "invalid": invalid,
                "reasons": reasons, "label": label, "verify_items": [], "message": "처리할 SN 이 없습니다"}
    ok_n, fails, dry = 0, [], None
    for i in range(0, len(devs), _CHUNK):
        res = S.register_devices(u["id"], devs[i:i + _CHUNK], allow_write=body.write)
        ok_n += res.get("success", 0)
        dry = res.get("dryRun") if dry is None else dry
        fails += [d for d in res.get("details", []) if not d.get("ok")]
    out = {"dry": bool(dry), "ok": not fails, "success": ok_n, "fails": fails[:30], "fail_count": len(fails),
           "invalid": invalid, "reasons": reasons, "label": label, "total": len(devs), "verify_items": []}
    if not dry:
        out["verify_items"] = devs
        if not fails:
            rec = S.record_registration(devices=devs, mail=_mail_ctx(g, u["id"]), processed_by=u["username"],
                                        category=cat, memo=f"[AI 해석 · {label}]")
            if rec.get("ok"):
                out["request"] = rec["request"]
        _invalidate()
    return out


@router.post("/organic/{tid}/current-dates")
def organic_current_dates(tid: str, request: Request):
    """등록일 변경 전 — 지금 FOTA 에 어떤 날짜로 들어가 있는지 샘플 10대(읽기)."""
    u = admin(request)
    from app.whitelist import fota as WF

    g = _find(tid)
    sns = _org_all(_org_ctx(g))
    sess = WF._session_from_user(u["id"])
    out = []
    for sn in sns[:10]:
        try:
            recs = WF.find_records(sess, sn) if sess else []
        except Exception:  # noqa: BLE001
            recs = []
        r0 = recs[0] if recs else {}
        out.append({"sn": sn, "date": r0.get("date", ""), "osType": r0.get("osType", "")})
    return {"total": len(sns), "rows": out, "login": sess is not None}


# ── DQA FOTA Test ─────────────────────────────────────────────────────────────

_dqa: dict[str, dict] = {}


def _dqa_work() -> Path:
    from app.database import BASE_DIR

    return BASE_DIR / "work" / "dqa_fota"


def _dqa_ctx(g: dict, uid: int, rebuild: bool = False) -> dict:
    from app.pipeline import dqa_fota as D
    from app.pipeline import dqa_jobs as DJ

    tid = _tid(g)
    ctx = _dqa.get(tid)
    if ctx is None or rebuild:
        rep = g["rep"]
        parsed = (ctx or {}).get("parsed") or D.parse_fota_test(
            rep.get("subject", "") or "", rep.get("body_text", "") or "", list(g.get("serials") or []))
        ov = (ctx or {}).get("os_type_override", "")
        ctx = {"parsed": parsed, "plan": D.build_plan(uid, parsed, os_type_override=ov),
               "osinfo_ok": False, "verify_items": [], "os_type_override": ov}
        _dqa[tid] = ctx
    job = DJ.snapshot(DJ.key_for(g["rep"]["id"]))
    if job:
        DJ.apply_to_plan(job, ctx["plan"])
    return ctx


def _dqa_view(g: dict, ctx: dict) -> dict:
    from app.pipeline import dqa_fota as D
    from app.pipeline import dqa_jobs as DJ

    p, plan = ctx["parsed"], ctx["plan"]
    job = DJ.snapshot(DJ.key_for(g["rep"]["id"]))
    jv = None
    if job:
        d = DJ.describe(job)
        jv = {"state": job["state"], "running": job["state"] == DJ.RUNNING, **d,
              "message": job.get("message") or "", "logs": list(job.get("logs") or [])[-200:],
              "restored": bool(job.get("restored"))}
    return {
        "parsed": {"product": p.product, "android": p.android, "sns": list(p.sns),
                   "froms": [{"dtype": r.dtype, "version": r.version, "path": r.path} for r in p.froms],
                   "to": ({"version": p.to.version, "path": p.to.path, "full_pfu": p.to.full_pfu,
                           "diff_pfu": p.to.diff_pfu} if p.to else None)},
        "plan": {"os_type": plan.os_type, "latest": plan.latest, "product": plan.product, "android": plan.android,
                 "api": plan.api, "end_date": plan.end_date, "note": plan.note,
                 "errors": list(plan.errors), "warnings": list(plan.warnings),
                 "version_prefix": plan.version_prefix, "version_prefix_src": plan.version_prefix_src,
                 "oi_rows": [{"osType": r.os_type, "version": r.version, "osVersion": r.os_version,
                              "description": r.description} for r in plan.oi_rows],
                 "dp_rows": [{"osVersion": r.os_version, "downloadType": r.dtype,
                              "fileNm": r.filename or "(다운로드 후 확정)", "size": r.size or "-",
                              "downloadUrl": r.download_url or ("(업로드 대기)" if not r.is_dummy else ""),
                              "dummy": r.is_dummy} for r in plan.dp_rows],
                 "uploaded": any(r.download_url for r in plan.dp_rows if not r.is_dummy)},
        "osinfo_ok": bool(ctx.get("osinfo_ok")),
        "verify_items": ctx.get("verify_items") or [],
        "whitelist_rows": _upload_rows(D.whitelist_devices(plan)),
        "job": jv,
    }


@router.get("/dqa/{tid}")
def dqa(tid: str, request: Request):
    u = user(request)
    g = _find(tid)
    return _dqa_view(g, _dqa_ctx(g, u["id"]))


class DqaFixIn(BaseModel):
    product: str = ""
    android: str = ""
    os_type: str = ""
    sns: list[str] = []
    froms: list[dict] = []
    to: dict | None = None


@router.post("/dqa/{tid}/recalc")
def dqa_recalc(tid: str, body: DqaFixIn, request: Request):
    u = user(request)
    from app.pipeline import dqa_fota as D

    g = _find(tid)
    ctx = _dqa_ctx(g, u["id"])
    p = ctx["parsed"]
    p.product = body.product.strip()
    pm = re.search(r"(\d{2,3})", p.product)
    p.product_num = pm.group(1) if pm else p.product_num
    p.android = body.android.strip()
    p.sns = [s.strip() for s in body.sns if s.strip()]
    for r, f in zip(p.froms, body.froms):
        r.version = (f.get("version") or "").strip()
        r.path = (f.get("path") or "").strip()
    if body.to and p.to:
        p.to.version = (body.to.get("version") or "").strip()
        p.to.path = (body.to.get("path") or "").strip()
        p.to.full_pfu = (body.to.get("full_pfu") or "").strip()
        p.to.diff_pfu = (body.to.get("diff_pfu") or "").strip()
    ctx["os_type_override"] = body.os_type.strip()
    ctx["plan"] = D.build_plan(u["id"], p, os_type_override=ctx["os_type_override"])
    ctx["osinfo_ok"] = False
    return _dqa_view(g, _dqa_ctx(g, u["id"]))


class ForceIn(BaseModel):
    force: bool = False


@router.post("/dqa/{tid}/download")
def dqa_download(tid: str, body: ForceIn, request: Request):
    u = admin(request)
    from app.pipeline import dqa_jobs as DJ

    g = _find(tid)
    ctx = _dqa_ctx(g, u["id"])
    key = DJ.key_for(g["rep"]["id"])
    if DJ.is_running(key):
        _bad("이미 진행 중입니다", 409)
    j = DJ.snapshot(key)
    if j and j["state"] == DJ.DONE and not body.force:
        return {"ok": False, "need_force": True,
                "message": "이 요청은 이미 다운로드와 Dropbox 업로드가 끝났습니다"}
    workdir = str(_dqa_work() / f"{g['rep']['id']}_{datetime.now():%Y%m%d-%H%M%S}")
    DJ.clear(key)
    ok, msg = DJ.start(key, u["id"], ctx["plan"], workdir, label=g["rep"].get("subject", ""))
    if not ok:
        _bad(msg, 409)
    return {"ok": True, "message": "백그라운드에서 시작했습니다"}


@router.post("/dqa/{tid}/cancel")
def dqa_cancel(tid: str, request: Request):
    admin(request)
    from app.pipeline import dqa_jobs as DJ

    g = _find(tid)
    return {"ok": DJ.cancel(DJ.key_for(g["rep"]["id"]))}


@router.post("/dqa/{tid}/clear")
def dqa_clear(tid: str, request: Request):
    admin(request)
    from app.pipeline import dqa_jobs as DJ

    g = _find(tid)
    DJ.clear(DJ.key_for(g["rep"]["id"]))
    return {"ok": True}


def _need_ostype(plan) -> None:
    if not plan.os_type:
        _bad("OS_Type 이 확정되지 않았습니다")


@router.post("/dqa/{tid}/fota-preview")
def dqa_fota_preview(tid: str, request: Request):
    u = user(request)
    from app.integrations import fota_system
    from app.pipeline import dqa_fota as D

    g = _find(tid)
    ctx = _dqa_ctx(g, u["id"])
    rep = D.register_fota(u["id"], ctx["plan"], False)
    return {"report": fota_system.format_report(rep)}


@router.post("/dqa/{tid}/osinfo")
def dqa_osinfo(tid: str, request: Request):
    u = admin(request)
    from app.integrations import fota_system
    from app.pipeline import dqa_fota as D

    g = _find(tid)
    ctx = _dqa_ctx(g, u["id"])
    _need_ostype(ctx["plan"])
    rr = D.register_osinfo(u["id"], ctx["plan"], True)
    ctx["osinfo_ok"] = bool(rr.get("osinfo_ok"))
    return {"ok": ctx["osinfo_ok"], "dry": rr.get("mode") == "DRY-RUN", "report": fota_system.format_report(rr)}


@router.post("/dqa/{tid}/osinfo-check")
def dqa_osinfo_check(tid: str, request: Request):
    u = user(request)
    from app.integrations import fota_system
    from app.pipeline import dqa_fota as D

    g = _find(tid)
    ctx = _dqa_ctx(g, u["id"])
    rr = fota_system.register_downloadpath_only(u["id"], D.osinfo_lists(ctx["plan"]), [], False)
    ctx["osinfo_ok"] = bool(rr.get("osinfo_ok"))
    return {"ok": ctx["osinfo_ok"], "report": "\n".join(rr.get("lines") or [])}


@router.post("/dqa/{tid}/downloadpath")
def dqa_downloadpath(tid: str, request: Request):
    u = admin(request)
    from app.integrations import fota_system
    from app.pipeline import dqa_fota as D

    g = _find(tid)
    ctx = _dqa_ctx(g, u["id"])
    plan = ctx["plan"]
    _need_ostype(plan)
    if not any(r.download_url for r in plan.dp_rows if not r.is_dummy):
        _bad("① 다운로드 + 업로드가 먼저입니다(to-OS 실 URL 필요)")
    rr = D.register_downloadpath(u["id"], plan, True)
    ctx["osinfo_ok"] = bool(rr.get("osinfo_ok"))
    return {"ok": bool(rr.get("ok")), "skipped": rr.get("skipped") or "",
            "dry": rr.get("mode") == "DRY-RUN", "report": fota_system.format_report(rr)}


@router.post("/dqa/{tid}/whitelist")
def dqa_whitelist(tid: str, request: Request):
    u = admin(request)
    from app.pipeline import dqa_fota as D

    S = _S()
    g = _find(tid)
    ctx = _dqa_ctx(g, u["id"])
    plan = ctx["plan"]
    if not plan.os_type or not ctx["parsed"].sns:
        _bad("OS_Type 또는 SN 이 없습니다")
    wl = D.register_whitelist(u["id"], plan, True)
    devs = D.whitelist_devices(plan)
    out = {"ok": bool(wl.get("ok")), "dry": bool(wl.get("dryRun")), "msg": wl.get("msg", ""),
           "details": [str(d) for d in wl.get("details", [])][:60]}
    if not wl.get("dryRun"):
        ctx["verify_items"] = devs
        if wl.get("ok"):
            rec = S.record_registration(devices=devs, mail=_mail_ctx(g, u["id"]), processed_by=u["username"],
                                        category="DQA", memo=f"[DQA FOTA Test · {plan.os_type} · api={plan.api}]")
            if rec.get("ok"):
                out["request"] = rec["request"]
        _invalidate()
    out["verify_items"] = ctx["verify_items"]
    return out


# ── 신청(사람이 직접) ───────────────────────────────────────────────────────────

@router.get("/options")
def options(request: Request):
    u = user(request)
    S = _S()
    from app.user_mailer import user_email

    rules = S.list_rules()
    return {"models": MODELS, "os_levels": OS_LEVELS, "categories": S.CATEGORIES,
            "default_category": S.DEFAULT_CATEGORY, "end_date": S.get_end_date(),
            "api": {str(k): v for k, v in S.OS_API.items()}, "rules": rules,
            "is_admin": u.get("role") == "admin", "email": user_email(u["id"]) or ""}


class BuildIn(BaseModel):
    sns: list[str]
    model: str = ""
    cur: str = ""
    upg: str = ""


@router.post("/build-devices")
def build_devices(body: BuildIn, request: Request):
    user(request)
    S = _S()
    return {"devices": [S.build_device(sn, body.cur, body.upg, body.model) for sn in body.sns if sn.strip()]}


@router.post("/bulk-parse")
async def bulk_parse(request: Request, file: UploadFile = File(...)):
    user(request)
    from openpyxl import load_workbook

    S = _S()
    data = await file.read()
    try:
        wb = load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001
        _bad(f"파일 읽기 오류: {e}")
    ws = wb.active
    out: list[dict] = []
    for row in ws.iter_rows(values_only=True):
        sn = str((row[0] if len(row) > 0 else "") or "").strip()
        if not sn:
            continue
        low = sn.lower()
        if low.startswith("sn") or "10자리" in sn or "serial" in low:
            continue
        if sn.startswith(("※", "▼")) or "WhiteList" in sn or "양식" in sn:
            continue
        out.append(S.build_device(sn, row[1] if len(row) > 1 else "", row[2] if len(row) > 2 else "",
                                  str((row[3] if len(row) > 3 else "") or "").strip()))
    wb.close()
    return {"devices": out}


class SubmitIn(BaseModel):
    devices: list[dict]
    charge_type: str = "무상"
    free_reason: str = ""
    memo: str = ""


@router.post("/requests")
def submit(body: SubmitIn, request: Request):
    u = user(request)
    from app.user_mailer import user_email
    from app.whitelist import mailer as MAIL

    S = _S()
    if not body.devices:
        _bad("등록할 단말기를 추가하세요")
    paid = body.charge_type == "유상"
    requester = {"username": u["username"], "email": user_email(u["id"]) or "", "name": u["username"], "team": ""}
    try:
        req = S.submit_request(requester=requester, devices=list(body.devices), memo=body.memo.strip(),
                               charge_type=body.charge_type,
                               free_reason="" if paid else body.free_reason.strip())
    except S.DuplicateSerialError as ex:
        lines = [f"{sn} — {v['request_id']} ({v['created_at']}, {STATUS_KO.get(v['status'], v['status'])})"
                 for sn, v in list(ex.duplicates.items())[:8]]
        raise HTTPException(409, "이미 신청된 단말이 있어 저장하지 않았습니다 · " + " · ".join(lines))
    try:
        mail_ok, mail_msg = MAIL.send_new_request(req, u["id"])
    except Exception as e:  # noqa: BLE001
        mail_ok, mail_msg = False, str(e)
    return {"request": req, "mail_ok": mail_ok, "mail_msg": mail_msg}


def _xlsx(build) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    build(wb, Font(bold=True, color="FFFFFF"), PatternFill("solid", fgColor="4A7C67"))
    bio = BytesIO()
    wb.save(bio)
    return bio.getvalue()


def _xlsx_resp(data: bytes, name: str) -> Response:
    from urllib.parse import quote

    return Response(data, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f"attachment; filename*=UTF-8''{quote(name)}"})


def _head(ws, hdr, font, fill, widths):
    ws.append(hdr)
    for c in ws[1]:
        c.font, c.fill = font, fill
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[chr(64 + i)].width = w


@router.get("/template.xlsx")
def template(request: Request):
    user(request)
    S = _S()

    def build(wb, font, fill):
        from collections import Counter

        ws = wb.active
        ws.title = "단말기_일괄등록"
        _head(ws, ["SN (10자리)", "Current OS", "Upgrade To", "Model"], font, fill, [22, 14, 14, 14])
        ws.append(["1234567890", 9, 11, "PM90"])
        ws.append(["2234567890", "", "", "PM45"])
        ws2 = wb.create_sheet("등록가능_규칙")
        _head(ws2, ["Model", "From OS", "To OS", "OS Type", "비고"], font, fill, [14, 12, 12, 22, 24])
        rules = S.list_rules()
        cnt = Counter(r["model"] for r in rules)
        for r in rules:
            ws2.append([r["model"], r["from"], r["to"], r["ostype"],
                        "단일 Rule — OS 생략 가능" if cnt[r["model"]] == 1 else ""])

    return _xlsx_resp(_xlsx(build), "WhiteList_일괄등록_양식.xlsx")


# ── 등록 내역 ──────────────────────────────────────────────────────────────────

def _filter_requests(q: str) -> list[dict]:
    allr = [r for r in _S().list_requests() if r["status"] != "deleted"]
    k = q.strip().lower()
    if k:
        allr = [r for r in allr if k in (r["userName"] or "").lower() or k in r["id"].lower()
                or any(k in d["sn"].lower() for d in r["devices"])]
    return allr


@router.get("/requests")
def requests_list(request: Request, q: str = "", limit: int = 200):
    u = user(request)
    from app.whitelist import portal as PT
    from app.whitelist.db import whitelist_session
    from app.whitelist.models import WLRequest

    total = len([r for r in _S().list_requests() if r["status"] != "deleted"])
    rows = _filter_requests(q)
    comps = PT.list_companies()
    cur = {}
    if comps:
        with whitelist_session() as db:
            cur = dict(db.query(WLRequest.id, WLRequest.company_id).filter(WLRequest.company_id.isnot(None)).all())
    out = []
    for r in rows[:max(1, min(limit, 1000))]:
        out.append({**r, "statusLabel": STATUS_KO.get(r["status"], r["status"]),
                    "poFileName": os.path.basename(r.get("poFilePath") or ""),
                    "companyId": cur.get(r["id"])})
    return {"total": total, "count": len(rows), "requests": out,
            "companies": [{"id": c["id"], "name": c["name"]} for c in comps],
            "is_admin": u.get("role") == "admin"}


@router.get("/requests/export.xlsx")
def requests_export(request: Request, q: str = ""):
    user(request)
    reqs = _filter_requests(q)

    def build(wb, font, fill):
        ws = wb.active
        ws.title = "Requests"
        _head(ws, ["Request ID", "요청자", "팀", "제출일", "SN", "Android API", "Support End Date", "OS Type",
                   "유/무상", "상태", "처리일"], font, fill, [18, 14, 10, 12, 16, 12, 16, 18, 8, 10, 12])
        for r in reqs:
            for d in r["devices"]:
                ws.append([r["id"], r["userName"], r["team"], str(r["createdAt"])[:10], d["sn"], d["api"],
                           d["endDate"], d["ostype"], r["chargeType"], STATUS_KO.get(r["status"], r["status"]),
                           str(r["processedAt"])[:10]])

    return _xlsx_resp(_xlsx(build), f"WhiteList_등록내역_{datetime.now():%Y%m%d}.xlsx")


@router.get("/requests/{req_id}/po")
def request_po(req_id: str, request: Request):
    user(request)
    r = next((x for x in _S().list_requests() if x["id"] == req_id), None)
    p = (r or {}).get("poFilePath") or ""
    if not p or not Path(p).is_file():
        raise HTTPException(404, "PO 파일이 없습니다")
    return FileResponse(p, filename=os.path.basename(p))


class CompanyIn(BaseModel):
    company_id: int | None = None


@router.post("/requests/{req_id}/company")
def request_company(req_id: str, body: CompanyIn, request: Request):
    admin(request)
    from app.whitelist import portal as PT

    ok, msg = PT.set_request_company(req_id, body.company_id or None)
    if not ok:
        _bad(msg)
    return {"ok": True, "message": msg}


class DeleteIn(BaseModel):
    reason: str = ""


@router.post("/requests/{req_id}/delete")
def request_delete(req_id: str, body: DeleteIn, request: Request):
    u = admin(request)
    ok, msg = _S().delete_request(req_id, body.reason.strip(), u["username"])
    if not ok:
        _bad(msg)
    return {"ok": True, "message": msg}


# ── 규칙 ───────────────────────────────────────────────────────────────────────

@router.get("/rules")
def rules(request: Request, cat: str = ""):
    user(request)
    S = _S()
    rs = S.list_rules(cat)
    return {"rules": [{**r, "api": S.get_api(r["from"])} for r in rs], "categories": S.CATEGORIES,
            "models": sorted(set(MODELS) | {r["model"] for r in S.list_rules()}), "os_levels": OS_LEVELS}


class RuleIn(BaseModel):
    category: str = "기타"
    model: str
    from_os: int
    to_os: int
    ostype: str


@router.post("/rules")
def rule_add(body: RuleIn, request: Request):
    admin(request)
    ok, msg = _S().add_rule(body.model.strip(), body.from_os, body.to_os, body.ostype.strip(), body.category)
    if not ok:
        _bad(msg)
    return {"ok": True}


@router.delete("/rules/{rule_id}")
def rule_delete(rule_id: int, request: Request):
    admin(request)
    if not _S().delete_rule(rule_id):
        raise HTTPException(404, "규칙이 없습니다")
    return {"ok": True}


@router.get("/rules/template.xlsx")
def rules_template(request: Request):
    user(request)
    S = _S()

    def build(wb, font, fill):
        ws = wb.active
        ws.title = "규칙"
        _head(ws, ["Category", "Model", "From OS", "To OS", "OS Type"], font, fill, [12, 12, 10, 10, 22])
        ws.append(["GMS", "PM95", 13, 15, "PM95GMS15"])
        ws.append(["DQA", "PM95", 13, 15, "WhiteDQA95V"])
        for r in S.list_rules():
            ws.append([r["category"], r["model"], r["from"], r["to"], r["ostype"]])

    return _xlsx_resp(_xlsx(build), "WhiteList_규칙_양식.xlsx")


@router.post("/rules/import")
async def rules_import(request: Request, file: UploadFile = File(...)):
    admin(request)
    from openpyxl import load_workbook

    data = await file.read()
    try:
        wb = load_workbook(BytesIO(data), data_only=True, read_only=True)
    except Exception as e:  # noqa: BLE001
        _bad(f"파일 읽기 오류: {e}")
    rows = []
    for row in wb.active.iter_rows(values_only=True):
        cat = str((row[0] if len(row) > 0 else "") or "").strip()
        model = str((row[1] if len(row) > 1 else "") or "").strip()
        if not model or model.lower() == "model" or cat.lower() == "category":
            continue
        rows.append({"category": cat, "model": model, "from": row[2] if len(row) > 2 else None,
                     "to": row[3] if len(row) > 3 else None,
                     "ostype": str((row[4] if len(row) > 4 else "") or "").strip()})
    wb.close()
    return _S().bulk_import_rules(rows)


# ── 고객 포털(고객사 · 계정) ─────────────────────────────────────────────────────

@router.get("/portal")
def portal(request: Request):
    admin(request)
    from app.whitelist import portal as PT

    accts = PT.list_accounts()
    for a in accts:
        a["status_label"] = PT.status_label(a["status"])
        a["last_login_at"] = str(a.get("last_login_at") or "")[:16]
    return {"companies": PT.list_companies(), "accounts": accts, "initial_password": PT.INITIAL_PASSWORD,
            "portal_url": PT.PORTAL_URL}


class CompIn(BaseModel):
    name: str | None = None
    email_domain: str | None = None
    active: bool | None = None


def _pt(res):
    ok, msg = res[0], res[1]
    if not ok:
        _bad(msg)
    return {"ok": True, "message": msg, **({"password": res[2]} if len(res) > 2 else {})}


@router.post("/portal/companies")
def comp_add(body: CompIn, request: Request):
    admin(request)
    from app.whitelist import portal as PT

    return _pt(PT.add_company(body.name or "", body.email_domain or ""))


@router.post("/portal/companies/{cid}")
def comp_update(cid: int, body: CompIn, request: Request):
    admin(request)
    from app.whitelist import portal as PT

    kw = {k: v for k, v in body.model_dump().items() if v is not None}
    return _pt(PT.update_company(cid, **kw))


@router.delete("/portal/companies/{cid}")
def comp_delete(cid: int, request: Request):
    admin(request)
    from app.whitelist import portal as PT

    return _pt(PT.delete_company(cid))


class AcctIn(BaseModel):
    username: str
    email: str = ""


@router.post("/portal/companies/{cid}/accounts")
def acct_add(cid: int, body: AcctIn, request: Request):
    admin(request)
    from app.whitelist import portal as PT

    return _pt(PT.add_account(cid, body.username, body.email))


@router.post("/portal/accounts/{aid}/reset")
def acct_reset(aid: int, request: Request):
    admin(request)
    from app.whitelist import portal as PT

    return _pt(PT.reset_password(aid))


class AcctStatusIn(BaseModel):
    status: str


@router.post("/portal/accounts/{aid}/status")
def acct_status(aid: int, body: AcctStatusIn, request: Request):
    admin(request)
    from app.whitelist import portal as PT

    return _pt(PT.set_account_status(aid, body.status))


@router.delete("/portal/accounts/{aid}")
def acct_delete(aid: int, request: Request):
    admin(request)
    from app.whitelist import portal as PT

    return _pt(PT.delete_account(aid))


# ── DQA 발신자(사내 인원 명부의 DQA 표시) ──────────────────────────────────────────

@router.get("/dqa-senders")
def dqa_senders(request: Request):
    admin(request)
    from app import staff

    S = _S()
    return {"staff": staff.list_staff(), "domains": S.dqa_senders().get("domains", [])}


class DqaFlagIn(BaseModel):
    is_dqa: bool


@router.post("/dqa-senders/{sid}")
def dqa_sender_set(sid: int, body: DqaFlagIn, request: Request):
    u = admin(request)
    from app import staff

    s = next((x for x in staff.list_staff() if x["id"] == sid), None)
    if not s:
        raise HTTPException(404, "대상을 찾을 수 없습니다")
    ok, msg = staff.save_staff(id=sid, name=s["name"], email=s["email"], status=s["status"], is_dqa=body.is_dqa,
                               is_handler=s["is_handler"], team=s["team"], memo=s["memo"], by=u["username"])
    if not ok:
        _bad(msg)
    _invalidate()
    return {"ok": True}
