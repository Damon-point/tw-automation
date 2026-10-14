"""
License 신청 메일 자동 인식 — 받은 메일에서 신청 건을 골라 신청서 칸을 미리 채운다.

신청은 대개 **사내 영업 담당자**가 보낸다("PMDMCE 36개 · 파트너사 ID allbc").
회사·수량·PMDM ID 는 본문에 있지만, 고객측 담당자(Contact)와 PO 는 없을 때가 많다.
그래서 찾은 값만 채우고, 비어 있는 칸은 사람이 넣는다. **지어내지 않는다.**

  1) 규칙 거르기 — 라이선스 낱말 + 요청 낱말이 함께 있고, Redmine·앱이 보낸 메일이 아닌 것
  2) Claude 해석 — 신청인지 판단하고 칸을 뽑는다. 본문 + 첨부(PDF·엑셀·워드) 글을 함께 읽는다
  3) 이미 등록된 신청과 맞춰 본다(PMDM ID + 회사) — 맞으면 '이미 신청됨' 으로 둔다

결과는 work/licenses/mail_inbox.json 에 메일별로 둔다(해석은 한 번만, 다시 읽기는 사람이 누른다).
"""
from __future__ import annotations

import json
import os
import re
import threading
from datetime import datetime, timedelta
from email.utils import parseaddr, parsedate_to_datetime
from pathlib import Path

from app.database import get_session
from app.license import service
from app.models import CachedMail

STORE = service.STORAGE_DIR / "mail_inbox.json"
WINDOW_DAYS = int(os.getenv("TWN_LICENSE_INBOX_DAYS", "45"))
MODEL = os.getenv("LICENSE_INTERPRET_MODEL", "claude-haiku-4-5")
PER_TICK = 6
PO_EXT = {".pdf", ".doc", ".docx", ".jpg", ".jpeg", ".png", ".xlsx", ".xls"}

_lock = threading.Lock()
_scan_lock = threading.Lock()

_LW = r"(?:licen[cs]es?|라이선스|라이센스|ライセンス)"
# 라이선스 낱말이 PMDM·Enterprise·키·발급·개수 바로 곁에 있어야 한다 — 'license' 가 지나가듯 나오는 메일(GMS 공지 등)은 거른다
_NEAR = re.compile(
    rf"(?:pmdm|enterprise|\bemm\b)[^\n]{{0,60}}{_LW}"
    rf"|{_LW}[^\n]{{0,40}}(?:\bkey\b|키|발급|추가|구매|요청|request)"
    rf"|\d[\d,]*\s*(?:개|ea|pcs|additional)?\s*{_LW}", re.I)
_NOT_SENDER = re.compile(r"pmissue@|redmine|tw\.auto@|no-?reply|atlassian|mailchimp", re.I)
_NOT_SUBJECT = re.compile(r"\[TW Automation\]|\[PMDM License|Issue #\d+", re.I)
_QUOTE = re.compile(r"^\s*(From:|보낸 사람:|差出人:|-----Original|-----Original Message|On .+ wrote:)", re.M)
_INLINE_IMG = re.compile(r"^image\d+\.(png|jpe?g|gif)$", re.I)

REQUIRED = [("company_name", "Company"), ("quantity", "Quantity"), ("contact_name", "Contact"),
            ("contact_email", "E-mail"), ("license_user_id", "PMDM ID")]


# ── 저장 ───────────────────────────────────────────────────────────────────────

def _load() -> dict:
    try:
        return json.loads(STORE.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return {}


def _save(d: dict) -> None:
    STORE.parent.mkdir(parents=True, exist_ok=True)
    tmp = STORE.with_suffix(".tmp")
    tmp.write_text(json.dumps(d, ensure_ascii=False, indent=1), encoding="utf-8")
    os.replace(tmp, STORE)


def _put(mail_id: int, **kv) -> dict:
    with _lock:
        d = _load()
        cur = d.get(str(mail_id), {})
        cur.update(kv)
        d[str(mail_id)] = cur
        _save(d)
        return cur


# ── 1) 규칙 거르기 ───────────────────────────────────────────────────────────────

def _top(body: str) -> str:
    """인용(예전 메일) 앞까지 — 신청 의도는 맨 위에 있다."""
    m = _QUOTE.search(body or "")
    return (body or "")[: m.start()] if m else (body or "")


def _when(m: CachedMail) -> datetime | None:
    try:
        return parsedate_to_datetime(m.received_at).replace(tzinfo=None)
    except Exception:  # noqa: BLE001
        return None


def looks_like_request(m: CachedMail) -> bool:
    if _NOT_SENDER.search(m.sender or "") or _NOT_SUBJECT.search(m.subject or ""):
        return False
    text = f"{m.subject or ''}\n{_top(m.body_text or '')[:4000]}"
    return bool(_NEAR.search(text))


_cand_cache: dict = {}


def candidates(days: int = WINDOW_DAYS) -> list[CachedMail]:
    """규칙에 걸린 후보 메일. 메일 판번호(건수·최대 id)가 그대로면 직전 결과를 쓴다 —
    본문 전체를 훑는 검색이라(메일 5천 통 · 0.6초) 화면을 열 때마다 하지 않는다."""
    from app.background import data_version

    key = (data_version(), days, datetime.now().strftime("%Y-%m-%d"))   # 날짜가 바뀌면 창도 바뀐다
    hit = _cand_cache.get("v")
    if hit and hit[0] == key:
        return list(hit[1])
    out = _candidates(days)
    _cand_cache["v"] = (key, out)
    return list(out)


def _candidates(days: int) -> list[CachedMail]:
    since = datetime.now() - timedelta(days=days)
    with get_session() as db:
        rows = (db.query(CachedMail)
                .filter(CachedMail.subject.ilike("%licen%") | CachedMail.subject.ilike("%라이선스%")
                        | CachedMail.subject.ilike("%라이센스%") | CachedMail.subject.ilike("%ライセンス%")
                        | CachedMail.body_text.ilike("%licen%") | CachedMail.body_text.ilike("%라이선스%")
                        | CachedMail.body_text.ilike("%라이센스%"))
                .order_by(CachedMail.id.desc()).limit(600).all())
        db.expunge_all()
    out = []
    for m in rows:
        w = _when(m)
        if w and w < since:
            continue
        if looks_like_request(m):
            out.append(m)
    return out


# ── 2) 첨부 글 ─────────────────────────────────────────────────────────────────

def attachments(m: CachedMail) -> list[str]:
    return [a for a in (m.attachment_names or "").splitlines()
            if a and not _INLINE_IMG.match(a) and os.path.splitext(a)[1].lower() in PO_EXT]


def _att_text(path: Path, limit: int = 3000) -> str:
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            import fitz

            with fitz.open(str(path)) as doc:
                return "\n".join(doc[i].get_text() for i in range(min(doc.page_count, 4)))[:limit]
        if ext in (".xlsx", ".xls"):
            import openpyxl

            wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets[:3]:
                for row in ws.iter_rows(max_row=60, values_only=True):
                    cells = [str(c) for c in row if c not in (None, "")]
                    if cells:
                        lines.append(" | ".join(cells))
            return "\n".join(lines)[:limit]
        if ext == ".docx":
            import docx

            d = docx.Document(str(path))
            parts = [p.text for p in d.paragraphs if p.text.strip()]
            for t in d.tables[:4]:
                for r in t.rows:
                    parts.append(" | ".join(c.text.strip() for c in r.cells))
            return "\n".join(parts)[:limit]
    except Exception:  # noqa: BLE001
        return ""
    return ""


def _att_texts(m: CachedMail) -> list[tuple[str, str]]:
    if not m.attachment_dir:
        return []
    out = []
    for a in attachments(m)[:4]:
        p = Path(m.attachment_dir) / os.path.basename(a)
        if p.is_file():
            t = _att_text(p)
            if t.strip():
                out.append((a, t))
    return out


# ── 3) Claude 해석 ──────────────────────────────────────────────────────────────

_SYS = """당신은 Point Mobile 기술지원팀에 온 메일을 읽고 **PMDM License 발급 신청**인지 판단해,
신청서 칸을 채웁니다. 신청은 대개 사내 영업 담당자(@pointmobile.com)가 고객/파트너 대신 보냅니다.

JSON 하나로만 답하세요:
{"is_request": true|false, "confidence": 0~1, "reason": "한국어 한 줄",
 "company_name": "", "quantity": 0, "license_user_id": "", "contact_name": "", "contact_email": "",
 "country": "", "requester": "", "po_number": "", "po_attachment": "", "sku": ""}

규칙:
· **판단 대상은 본문 맨 위(가장 최근) 메시지다.** 아래 인용된 예전 메일은 칸을 채울 정보로만 쓴다.
  맨 위가 회신·확인·발급 완료 안내·내부 컨펌 요청이면 is_request=false.
· is_request: 라이선스(키) 발급·추가를 **요청**하는 메일만 true. 발급 결과 안내, 내부 보고/컨펌 요청,
  제품 이슈·기능 문의, 가격 문의, Redmine 알림은 false.
· company_name: 라이선스를 받을 고객사(또는 파트너사) 정식 이름. PO·인보이스에 정식 상호가 있으면 그것.
· quantity: 라이선스 개수(정수). "36개", "50 licenses", "23EA" 등.
· license_user_id: PMDM(Enterprise) 계정 ID. "Account ID: mobilscan(335)" 이면 "mobilscan".
· contact_name / contact_email: **고객측** 담당자. PO·인보이스·본문에 고객 담당자가 있을 때만.
  Point Mobile 직원(@pointmobile.com)은 넣지 마세요.
· requester: 요청한 사내 담당자 "이름(팀)" — 예 "김한울(해외영업 아시아팀)".
· po_number: PO/인보이스 번호가 보이면. po_attachment: 첨부 중 PO/인보이스 파일 이름(첨부 목록에 있는 그대로).
· sku: PMDMCE 같은 제품 코드가 보이면.
· **모르면 빈 문자열(수량은 0)**. 추측해서 채우지 마세요."""


def _prompt(m: CachedMail, atts: list[tuple[str, str]]) -> str:
    body = (m.body_text or "").strip()
    parts = [f"제목: {m.subject}", f"보낸 사람: {m.sender}", f"받은 때: {m.received_at}",
             f"첨부: {', '.join(attachments(m)) or '(없음)'}", "", "본문:", body[:6000]]
    for name, t in atts:
        parts += ["", f"--- 첨부 {name} 의 글", t]
    return "\n".join(parts)


def _ai(m: CachedMail) -> dict:
    from app.integrations.claude_cli import have_auth

    if not have_auth():
        return {"error": "Claude 를 부를 수 없습니다"}
    import anthropic

    from app import usage
    from app.pipeline.ai_edit import _extract_json

    usage.set_operation("License 신청 메일 해석", m.id)
    resp = anthropic.Anthropic().messages.create(
        model=MODEL, max_tokens=800, system=_SYS,
        messages=[{"role": "user", "content": _prompt(m, _att_texts(m))}])
    text = "".join(b.text for b in resp.content if getattr(b, "type", None) == "text")
    p = _extract_json(text) or {}
    try:
        q = int(str(p.get("quantity") or 0).replace(",", "").strip() or 0)
    except ValueError:
        q = 0
    s = lambda k, n=200: str(p.get(k) or "").strip()[:n]  # noqa: E731
    email = s("contact_email")
    if email.lower().endswith("@pointmobile.com"):
        email = ""
    return {"is_request": bool(p.get("is_request")), "confidence": float(p.get("confidence") or 0),
            "reason": s("reason", 300), "company_name": s("company_name"), "quantity": q,
            "license_user_id": s("license_user_id", 120), "contact_name": s("contact_name", 120),
            "contact_email": email, "country": s("country", 60), "requester": s("requester", 120),
            "po_number": s("po_number", 80), "po_attachment": s("po_attachment", 300), "sku": s("sku", 40)}


# ── 이미 등록된 신청과 맞추기 ─────────────────────────────────────────────────────

def _norm(s: str) -> str:
    return re.sub(r"[^a-z0-9가-힣]", "", (s or "").lower())


def _match_existing(ai: dict, when: datetime | None) -> str:
    uid, comp = _norm(ai.get("license_user_id", "")), _norm(ai.get("company_name", ""))
    if not uid and not comp:
        return ""
    for r in service.list_all():
        if when and r.created_at and r.created_at < when - timedelta(days=2):
            continue
        if uid and _norm(r.license_user_id) == uid and (not ai.get("quantity") or
                                                         service.norm_qty(r.quantity) == str(ai["quantity"])):
            return r.doc_id
        if comp and uid and _norm(r.company_name)[:8] == comp[:8] and _norm(r.license_user_id) == uid:
            return r.doc_id
    return ""


# ── 우리 팀이 이미 회신한 건 ─────────────────────────────────────────────────────

_PREFIX = re.compile(r"^\s*((re|fw|fwd|aw|wg|sv|回覆|回复|答复|转发|転送)\s*[:：]\s*)+", re.I)


def _thread_key(subject: str) -> str:
    return _PREFIX.sub("", subject or "").strip()


def _team() -> set[str]:
    from app import user_mailer
    from app.models import User

    with get_session() as db:
        ids = [u.id for u in db.query(User).all()]
    return {e.lower() for e in (user_mailer.user_email(i) for i in ids) if e}


def answered_by(mail_id: int, subject: str, team: set[str] | None = None) -> str:
    """같은 제목 줄기에서 이 메일 뒤에 우리 팀(앱 사용자)이 보낸 메일이 있으면 그 사람 주소.

    PMDM ID 가 메일에 없어 기존 신청과 못 맞춘 건도, 우리가 이미 답했다면 처리된 것으로 본다.
    """
    key = _thread_key(subject)
    if len(key) < 6:
        return ""
    team = team if team is not None else _team()
    with get_session() as db:
        rows = (db.query(CachedMail.sender, CachedMail.subject)
                .filter(CachedMail.id > mail_id, CachedMail.subject.like(f"%{key}%")).all())
    for sender, subj in rows:
        addr = parseaddr(sender or "")[1].lower()
        if addr in team and _thread_key(subj) == key:
            return addr
    return ""


# ── 훑기 ───────────────────────────────────────────────────────────────────────

def interpret(mail_id: int) -> dict:
    with get_session() as db:
        m = db.get(CachedMail, mail_id)
        if m is None:
            raise KeyError(mail_id)
        db.expunge(m)
    try:
        ai = _ai(m)
    except Exception as e:  # noqa: BLE001
        ai = {"error": f"{type(e).__name__}: {e}"[:300]}
    rec = {"ai": ai, "subject": m.subject or "", "scanned_at": datetime.now().isoformat(timespec="seconds")}
    cur = _load().get(str(mail_id), {})
    if not cur.get("status") or cur.get("status") == "new":
        linked = _match_existing(ai, _when(m)) if ai.get("is_request") else ""
        rec["status"] = "linked" if linked else "new"
        if linked:
            rec["doc_id"] = linked
    return _put(mail_id, **rec)


def scan(limit: int = PER_TICK) -> int:
    """새 후보만 해석한다. 동시에 두 번 돌지 않는다."""
    if not _scan_lock.acquire(blocking=False):
        return 0
    try:
        known = _load()
        todo = [m.id for m in candidates() if str(m.id) not in known]
        for mid in todo[:limit]:
            interpret(mid)
        return len(todo[:limit])
    finally:
        _scan_lock.release()


def scan_async() -> None:
    threading.Thread(target=scan, name="twn-license-scan", daemon=True).start()


# ── 화면용 ─────────────────────────────────────────────────────────────────────

def _fields(ai: dict) -> dict:
    return {k: ai.get(k) or ("" if k != "quantity" else None) for k, _ in REQUIRED}


def missing(ai: dict) -> list[str]:
    return [lab for k, lab in REQUIRED if not ai.get(k)]


def _note(ai: dict, subject: str) -> str:
    bits = []
    if ai.get("requester"):
        bits.append(f"{ai['requester']} 요청 건")
    extra = " / ".join(x for x in (ai.get("country"), ai.get("sku"), ai.get("po_number")) if x)
    if extra:
        bits.append(extra)
    bits.append(f"메일: {subject}")
    return " · ".join(bits)[:500]


def item(m: CachedMail, rec: dict) -> dict:
    ai = rec.get("ai") or {}
    atts = attachments(m)
    po = ai.get("po_attachment") if ai.get("po_attachment") in atts else ""
    if not po:
        po = next((a for a in atts if re.search(r"\bpo\b|purchase|order|invoice|인보이스|발주", a, re.I)), "")
    name, addr = parseaddr(m.sender or "")
    return {
        "mail_id": m.id, "subject": m.subject or "", "sender": (name or addr).strip('" '), "sender_email": addr,
        "received_at": m.received_at, "status": rec.get("status", "pending"), "doc_id": rec.get("doc_id", ""),
        "scanned": bool(ai), "error": ai.get("error", ""),
        "is_request": bool(ai.get("is_request")), "confidence": ai.get("confidence", 0), "reason": ai.get("reason", ""),
        "fields": _fields(ai), "missing": missing(ai) + ([] if po else ["PO 파일"]),
        "requester": ai.get("requester", ""), "country": ai.get("country", ""), "sku": ai.get("sku", ""),
        "po_number": ai.get("po_number", ""), "note": _note(ai, m.subject or ""),
        "attachments": atts, "po_guess": po,
    }


def inbox() -> dict:
    d = _load()
    cands = candidates()
    ids = {m.id for m in cands}
    # 예전에 해석해 둔 것 중 창 밖으로 나간 것도, 사람이 처리 중이면 보여 준다
    extra = [int(k) for k, v in d.items() if int(k) not in ids and v.get("status") == "new"]
    with get_session() as db:
        more = db.query(CachedMail).filter(CachedMail.id.in_(extra)).all() if extra else []
        db.expunge_all()
    items, pending = [], 0
    team = _team()
    for m in sorted(cands + more, key=lambda x: x.id, reverse=True):
        rec = d.get(str(m.id))
        if rec is None:
            pending += 1
            continue
        ai = rec.get("ai") or {}
        if not ai.get("is_request") and rec.get("status") not in ("created",):
            continue
        it = item(m, rec)
        if it["status"] == "new":
            by = answered_by(m.id, m.subject or "", team)
            if by:
                it["status"], it["answered_by"] = "answered", by
        items.append(it)
    return {"items": items, "pending": pending,
            "counts": {s: sum(1 for i in items if i["status"] == s)
                       for s in ("new", "answered", "linked", "created", "dismissed")}}


def detail(mail_id: int) -> dict:
    with get_session() as db:
        m = db.get(CachedMail, mail_id)
        if m is None:
            raise KeyError(mail_id)
        db.expunge(m)
    rec = _load().get(str(mail_id)) or interpret(mail_id)
    out = item(m, rec)
    out["body"] = (m.body_text or "")[:8000]
    return out


def set_status(mail_id: int, status: str, doc_id: str = "", by: str = "") -> dict:
    return _put(mail_id, status=status, doc_id=doc_id, decided_by=by,
                decided_at=datetime.now().isoformat(timespec="seconds"))


def open_count() -> int:
    """아직 아무도 손대지 않은 신청 수 — 사이드바 숫자."""
    team = None
    n = 0
    for k, v in _load().items():
        if v.get("status") != "new" or not (v.get("ai") or {}).get("is_request"):
            continue
        team = team if team is not None else _team()
        if not answered_by(int(k), v.get("subject", ""), team):
            n += 1
    return n
