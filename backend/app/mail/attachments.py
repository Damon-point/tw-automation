"""메일 첨부 처리 — WhiteList 요청은 SN 을 **여러 곳에 흩어서** 보낸다.

실측(‘[유럽팀] 등록일 변경’): 본문 표 + 엑셀 첨부(ALDI_SWAP) + **첨부된 예전 메일들
(message/rfc822, 파일명 없음)** 안에 제품·OS 가 제각각인 SN 들이 텍스트로 들어 있다.
"요청에 틀이 없다"는 게 이 구조다. 그래서 여기서는 **재귀적으로** 훑는다:

  · 엑셀/CSV 첨부      → 모든 셀에서 SN 후보
  · message/rfc822 / .eml 첨부(예전 메일) → 그 **본문 텍스트 + 그 안의 엑셀**까지 재귀
  · .msg(Outlook)     → extract_msg 있으면 읽고, 없으면 이름만 기록

등록일 변경은 SN 을 넉넉히 모아도 안전하다 — DB 반영이 **등록된 SN 만**
날짜를 바꾸고 나머지는 '미등록'으로 건너뛴다. 첨부 파일은 work/mail_attachments/ 에 보관.
"""
from __future__ import annotations

import csv
import email
import io
import re
from email.header import decode_header, make_header
from email.message import Message
from pathlib import Path

# SN 스캔 규칙은 본문과 동일하게 맞춘다(중복 정의 없이 재사용).
from app.whitelist.mail_filter import is_serial_like

_ATTACH_ROOT = Path(__file__).resolve().parent.parent.parent / "work" / "mail_attachments"

_SN_EXTS = (".xlsx", ".xlsm", ".xls", ".csv")    # SN 추출 대상
_KEEP_EXTS = (".xlsx", ".xlsm", ".xls", ".csv")  # 파일 보관 대상
_TOKEN_RE = re.compile(r"(?<![A-Za-z0-9])([A-Za-z0-9]{10})(?![A-Za-z0-9])")
_XLRD_WARNED = False                             # 없다는 경고는 한 번만


def _log(msg: str) -> None:
    """dev.log/server.log 로 나가는 표준출력에 남긴다(background._log 와 같은 형식)."""
    from datetime import datetime

    print(f"[attach {datetime.now():%Y-%m-%d %H:%M:%S}] {msg}", flush=True)
_TAG_RE = re.compile(r"<[^>]+>")
_MAX_ATTACH_BYTES = 25 * 1024 * 1024
_MAX_DEPTH = 6


def _safe(name: str) -> str:
    return re.sub(r"[^A-Za-z0-9._-]", "_", (name or "").strip()) or "file"


def _dec(value: str | None) -> str:
    if not value:
        return ""
    try:
        return str(make_header(decode_header(value)))
    except Exception:  # noqa: BLE001
        return value


def _payload_bytes(part: Message) -> bytes | None:
    try:
        data = part.get_payload(decode=True)
    except Exception:  # noqa: BLE001
        return None
    if not data or len(data) > _MAX_ATTACH_BYTES:
        return None
    return data


def _payload_text(part: Message) -> str:
    data = part.get_payload(decode=True)
    if data is None:
        return ""
    charset = part.get_content_charset() or "utf-8"
    try:
        return data.decode(charset, errors="replace")
    except (LookupError, TypeError):
        return data.decode("utf-8", errors="replace")


def _serials_from_text(text: str) -> list[str]:
    out, seen = [], set()
    for m in _TOKEN_RE.finditer(text or ""):
        tok = m.group(1)
        if is_serial_like(tok):
            up = tok.upper()
            if up not in seen:
                seen.add(up)
                out.append(up)
    return out


def _serials_from_cells(cells) -> list[str]:
    return _serials_from_text(" ".join("" if c is None else str(c) for c in cells))


def _serials_from_xls(data: bytes) -> list[str]:
    """구형 엑셀(.xls, BIFF) — xlrd 로 읽는다. openpyxl 은 이 포맷을 못 연다.

    xlrd 가 없으면 빈 목록. 설치되지 않은 환경에서 메일 수신 자체가 죽으면 안 된다.
    실측: `.xls` 가 엑셀 첨부의 42%(13/31)였고 전부 SN 추출이 0건이었다.
    """
    try:
        import xlrd  # type: ignore
    except Exception:  # noqa: BLE001
        global _XLRD_WARNED
        if not _XLRD_WARNED:                  # 매 첨부마다 같은 줄을 찍지 않는다
            _XLRD_WARNED = True
            _log("xlrd 가 없어 .xls 첨부의 SN 을 읽지 못합니다 — pip install xlrd")
        return []
    book = xlrd.open_workbook(file_contents=data)
    cells = []
    for sh in book.sheets():
        for r in range(sh.nrows):
            # xlrd 는 숫자 셀을 **항상 float** 로 준다(openpyxl 은 int). 그대로 두면
            # SN 이 '2232710498.0' 이 돼 10자리 토큰 규칙에 걸리지 않을 수 있다.
            cells.extend(
                int(c) if isinstance(c, float) and c.is_integer() else c
                for c in sh.row_values(r)
            )
    return _serials_from_cells(cells)


def serials_from_attachment(filename: str, data: bytes) -> list[str]:
    """엑셀/CSV 첨부의 모든 셀에서 SN 후보를 뽑는다.

    실패해도 예외를 올리지 않는다(메일 한 통 때문에 수신이 멈추면 안 된다). 대신
    **반드시 로그를 남긴다** — 안 그러면 '첨부가 깨져서 0건'과 'SN 이 원래 없어서
    0건'을 구분할 수 없다(실측: .xls 13건이 전부 0건이었는데 파서가 아예 없었다).
    """
    name = (filename or "").lower()
    try:
        if name.endswith(".xls"):
            return _serials_from_xls(data)
        if name.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            cells = []
            try:
                for ws in wb.worksheets:
                    for row in ws.iter_rows(values_only=True):
                        cells.extend(row)
            finally:
                wb.close()
            return _serials_from_cells(cells)
        if name.endswith(".csv"):
            text = data.decode("utf-8-sig", errors="replace")
            cells = []
            for row in csv.reader(io.StringIO(text)):
                cells.extend(row)
            return _serials_from_cells(cells)
    except Exception as e:  # noqa: BLE001
        _log(f"첨부 파싱 실패 — {filename} ({len(data):,}B) · "
             f"{type(e).__name__}: {str(e)[:160]}")
        return []
    return []


def _serials_from_msg(data: bytes) -> tuple[list[str], bool]:
    """Outlook .msg — extract_msg 가 있으면 본문·첨부에서 SN 을 뽑는다. (serials, 처리됨)."""
    try:
        import extract_msg  # type: ignore
    except Exception:  # noqa: BLE001
        return [], False
    try:
        m = extract_msg.openMsg(io.BytesIO(data))
        out = list(_serials_from_text(m.body or ""))
        for att in getattr(m, "attachments", []) or []:
            fn = getattr(att, "longFilename", "") or getattr(att, "shortFilename", "") or ""
            ad = getattr(att, "data", None)
            if isinstance(ad, bytes) and fn.lower().endswith(_SN_EXTS):
                out += serials_from_attachment(fn, ad)
        return out, True
    except Exception as e:  # noqa: BLE001
        _log(f".msg 첨부 파싱 실패 — {type(e).__name__}: {str(e)[:160]}")
        return [], True


class _Acc:
    def __init__(self):
        self.serials: list[str] = []
        self.seen: set[str] = set()
        self.names: list[str] = []
        self.keep: list[tuple[str, bytes]] = []

    def add(self, sns):
        for s in sns:
            u = s.upper()
            if u not in self.seen:
                self.seen.add(u)
                self.serials.append(u)


def _collect(part: Message, acc: _Acc, *, is_nested: bool, depth: int) -> None:
    if depth > _MAX_DEPTH:
        return
    ct = part.get_content_type()

    # 첨부된 예전 메일 — 본문 텍스트 + 그 안의 엑셀까지 재귀.
    if ct == "message/rfc822":
        payload = part.get_payload()
        sub = payload[0] if isinstance(payload, list) and payload else payload
        if isinstance(sub, Message):
            acc.names.append("[첨부메일] " + _dec(sub.get("Subject", "")))
            _collect(sub, acc, is_nested=True, depth=depth + 1)
        return

    # 컨테이너(멀티파트)면 자식으로 내려간다.
    if part.is_multipart():
        for p in part.get_payload():
            _collect(p, acc, is_nested=is_nested, depth=depth + 1)
        return

    fname = _dec(part.get_filename() or "")
    low = fname.lower()
    disp = str(part.get("Content-Disposition") or "").lower()

    # .eml / .msg 첨부(예전 메일)
    if low.endswith(".eml"):
        data = _payload_bytes(part)
        if data:
            sub = email.message_from_bytes(data)
            acc.names.append("[첨부메일] " + _dec(sub.get("Subject", "")))
            _collect(sub, acc, is_nested=True, depth=depth + 1)
        return
    if low.endswith(".msg"):
        data = _payload_bytes(part)
        sns, ok = _serials_from_msg(data or b"")
        acc.names.append("[Outlook] " + fname + ("" if ok else " (extract-msg 미설치)"))
        acc.add(sns)
        return

    # 엑셀/CSV 첨부 — SN 추출 + 파일 보관
    if low.endswith(_KEEP_EXTS):
        data = _payload_bytes(part)
        if data:
            acc.add(serials_from_attachment(fname, data))
            acc.names.append(fname)
            acc.keep.append((fname, data))
        return

    # 텍스트 본문 — **첨부메일 내부에서만** SN 스캔(최상위 본문은 classify 가 처리).
    if is_nested and ct in ("text/plain", "text/html") and "attachment" not in disp:
        txt = _payload_text(part)
        if ct == "text/html":
            txt = _TAG_RE.sub(" ", txt)
        acc.add(_serials_from_text(txt))


_IMG_TYPES = {"image/png", "image/jpeg", "image/jpg", "image/gif", "image/webp"}
_IMG_MIN = 3 * 1024            # 이보다 작으면 서명/스페이서 아이콘일 확률↑ (건너뜀)
_IMG_MAX = 5 * 1024 * 1024
_IMG_MAX_COUNT = 6


def collect_body_images(msg: Message) -> list[dict]:
    """**본문**의 인라인 이미지(표·스크린샷)를 모은다 — OS 버전 등이 이미지로 오는 경우.

    첨부메일(message/rfc822) 안의 이미지(주로 서명 로고)는 제외한다. Claude 멀티모달로
    바로 읽히도록 {media_type, data(bytes)} 리스트를 돌려준다(개수·크기 제한).
    """
    out: list[dict] = []

    def _walk(part: Message, depth: int = 0) -> None:
        if depth > 5 or len(out) >= _IMG_MAX_COUNT:
            return
        ct = part.get_content_type()
        if ct == "message/rfc822":     # 첨부메일 안(서명 로고 등)은 보지 않는다
            return
        if part.is_multipart():
            for p in part.get_payload():
                _walk(p, depth + 1)
            return
        mt = (ct or "").lower()
        if mt == "image/jpg":
            mt = "image/jpeg"
        if mt in _IMG_TYPES:
            data = part.get_payload(decode=True)
            if data and _IMG_MIN <= len(data) <= _IMG_MAX:
                out.append({"media_type": mt, "data": data})

    _walk(msg)
    return out


def save_attachments(message_id: str, files: list[tuple[str, bytes]]) -> str:
    if not files:
        return ""
    d = _ATTACH_ROOT / _safe(message_id or "nomsgid")[:120]
    d.mkdir(parents=True, exist_ok=True)
    for fname, data in files:
        try:
            (d / _safe(fname)).write_bytes(data)
        except Exception:  # noqa: BLE001
            continue
    return str(d)


def process_message(msg: Message, message_id: str, raw: bytes | None = None) -> dict:
    """메시지의 첨부(엑셀·CSV·첨부메일)를 재귀적으로 훑어 SN 을 모으고 파일을 보관한다.

    첨부/첨부메일/본문이미지가 있으면 **원문(.eml)도 폴더에 저장**한다 — 나중에 처리 팝업이
    POP3 재조회 없이 로컬에서 블록·이미지·AI 해석을 하도록(동시 접속 1개 제약 회피).
    반환 {names, serials, dir}.
    """
    acc = _Acc()
    _collect(msg, acc, is_nested=False, depth=0)
    has_media = bool(acc.keep or acc.names or collect_body_images(msg))
    dir_path = ""
    if has_media:
        dir_path = save_attachments(message_id, acc.keep)  # 폴더 생성 + 엑셀 보관
        if not dir_path:
            dir_path = str(_ATTACH_ROOT / _safe(message_id or "nomsgid")[:120])
            Path(dir_path).mkdir(parents=True, exist_ok=True)
        if raw:
            try:
                (Path(dir_path) / RAW_EML).write_bytes(raw)
            except Exception:  # noqa: BLE001
                pass
    return {"names": acc.names, "serials": acc.serials, "dir": dir_path}


RAW_EML = "_raw.eml"


def load_raw(dir_path: str) -> bytes | None:
    """저장해 둔 원문(.eml)을 읽는다. 없으면 None."""
    if not dir_path:
        return None
    p = Path(dir_path) / RAW_EML
    try:
        return p.read_bytes() if p.exists() else None
    except Exception:  # noqa: BLE001
        return None


_PLAN_JSON = "_plan.json"


def save_plan(dir_path: str, plan: list) -> None:
    """AI 해석 결과를 캐시(폴더/_plan.json). 다이얼로그 재오픈 시 재과금 방지."""
    if not dir_path:
        return
    import json
    try:
        (Path(dir_path) / _PLAN_JSON).write_text(
            json.dumps(plan, ensure_ascii=False), encoding="utf-8")
    except Exception:  # noqa: BLE001
        pass


def load_plan(dir_path: str) -> list | None:
    if not dir_path:
        return None
    import json
    p = Path(dir_path) / _PLAN_JSON
    try:
        return json.loads(p.read_text(encoding="utf-8")) if p.exists() else None
    except Exception:  # noqa: BLE001
        return None


# ── 소스 블록 추출(유기적 처리용) ────────────────────────────────────────────
# 요청을 '본문 / 첨부메일 각각 / 엑셀 각각' 으로 나눈다. 각 블록의 SN 은 결정론적으로 뽑고,
# 블록의 텍스트는 Claude 해석(유형·모델·OS·날짜)의 근거로 쓴다.
def excel_text_sample(filename: str, data: bytes, max_rows: int = 25) -> str:
    """엑셀/CSV 앞부분을 텍스트로(모델·OS 열을 Claude 가 읽도록). SN 자체는 따로 뽑는다."""
    name = (filename or "").lower()
    rows_out: list[str] = []
    try:
        if name.endswith((".xlsx", ".xlsm")):
            from openpyxl import load_workbook
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            try:
                ws = wb.worksheets[0]
                for i, row in enumerate(ws.iter_rows(values_only=True)):
                    if i >= max_rows:
                        break
                    rows_out.append(" | ".join("" if c is None else str(c) for c in row))
            finally:
                wb.close()
        elif name.endswith(".csv"):
            text = data.decode("utf-8-sig", errors="replace")
            for i, row in enumerate(csv.reader(io.StringIO(text))):
                if i >= max_rows:
                    break
                rows_out.append(" | ".join(row))
    except Exception:  # noqa: BLE001
        return ""
    return "\n".join(rows_out)


def _own_text(msg: Message) -> str:
    """이 메시지 **자신의** 본문 텍스트(첨부메일/첨부는 내려가지 않는다)."""
    parts: list[str] = []

    def _t(part: Message, depth: int = 0) -> None:
        if depth > 4:
            return
        ct = part.get_content_type()
        if ct == "message/rfc822":
            return
        disp = str(part.get("Content-Disposition") or "").lower()
        if part.is_multipart():
            for p in part.get_payload():
                _t(p, depth + 1)
            return
        if ct in ("text/plain", "text/html") and "attachment" not in disp \
                and not part.get_filename():
            txt = _payload_text(part)
            if ct == "text/html":
                txt = _TAG_RE.sub(" ", txt)
            parts.append(txt)

    _t(msg)
    seen, uniq = set(), []
    for t in parts:
        k = " ".join((t or "").split())[:200]
        if k and k not in seen:
            seen.add(k)
            uniq.append(t)
    return "\n".join(uniq).strip()


def extract_blocks(msg: Message, *, depth: int = 0, label: str = "본문",
                   out: list | None = None) -> list[dict]:
    """요청을 소스 블록으로 나눈다. 각 블록 = {source, text, serials}.

    본문 · 첨부메일(재귀) · 엑셀/CSV 각각을 별도 블록으로 만든다.
    """
    out = out if out is not None else []
    if depth > _MAX_DEPTH:
        return out
    text = _own_text(msg)
    out.append({"source": label, "text": text, "serials": _serials_from_text(text)})

    def _scan(part: Message) -> None:
        ct = part.get_content_type()
        if ct == "message/rfc822":
            payload = part.get_payload()
            sub = payload[0] if isinstance(payload, list) and payload else payload
            if isinstance(sub, Message):
                extract_blocks(sub, depth=depth + 1,
                               label="[첨부메일] " + _dec(sub.get("Subject", "")), out=out)
            return
        if part.is_multipart():
            for p in part.get_payload():
                _scan(p)
            return
        fn = _dec(part.get_filename() or "")
        low = fn.lower()
        if low.endswith(".eml"):
            data = _payload_bytes(part)
            if data:
                sub = email.message_from_bytes(data)
                extract_blocks(sub, depth=depth + 1,
                               label="[첨부메일] " + _dec(sub.get("Subject", "")), out=out)
            return
        if low.endswith(_SN_EXTS):
            data = _payload_bytes(part)
            if data:
                out.append({"source": f"[엑셀] {fn}",
                            "text": excel_text_sample(fn, data),
                            "serials": serials_from_attachment(fn, data)})

    if msg.is_multipart():
        for p in msg.get_payload():
            _scan(p)
    return out
