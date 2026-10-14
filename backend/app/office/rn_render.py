"""RN 렌더링 — Word 템플릿 채우기(docxtpl) → PDF 변환(docx2pdf/Word COM), Excel 템플릿 채우기.

- render_rn_docx(context): rn_docx 템플릿을 docxtpl(Jinja2)로 채워 docx 바이트 반환.
- docx_to_pdf_bytes(docx_bytes): Word COM 으로 PDF 변환(윈도우+Word 필요). 워커 스레드 COM 초기화 처리.
- render_rn_pdf(context): 위 둘을 이어 PDF 바이트 반환.
- fill_excel_template(rows): rn_excel 템플릿을 열어 이슈 행을 채워 xlsx 바이트 반환.

템플릿의 정확한 placeholder/열 매핑은 사용자가 템플릿을 주면 확정한다(지금은 범용 골격).
로컬 전용.
"""
from __future__ import annotations

import io
import os
import tempfile
import threading
import time

from app.office import templates as T


def render_rn_docx(context: dict, template_key: str = "rn_docx") -> tuple[bytes | None, str]:
    """Word 템플릿을 context 로 채워 docx 바이트 반환. (bytes, error).

    template_key: 'rn_docx'(RN for Sales · 한 판) 또는 'rn_docx_full'(전체 이력 · 여러 판).

    autoescape=True 필수 — 끄면 값의 `&`·`<` 가 XML 에 **날것으로** 들어가 문서가 깨진다.
    실측: 'Added backup & restore items' → '&' 소실, 'Fixed <b> tag' → 그 줄부터 **뒤가 통째로
    사라짐**(lxml 이 조용히 복구하며 버린다). RN 문구에 '&' 는 실제로 들어온다.
    """
    path = T.path_for(template_key)
    if not path or not path.exists():
        label = (T.TEMPLATES.get(template_key) or ("", template_key, ""))[1]
        return None, f"Word 템플릿({template_key})이 없습니다 — {label}. [템플릿 관리]에서 업로드하세요."
    try:
        from docxtpl import DocxTemplate

        doc = DocxTemplate(str(path))
        doc.render(context, autoescape=True)
        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue(), ""
    except Exception as e:  # noqa: BLE001
        return None, f"Word 템플릿 렌더 실패: {e}"


#: Word 는 자동화에서 스레드 안전하지 않고, docx2pdf 는 한 번 변환할 때마다 Word 를
#: 띄웠다 닫는다. 그래서 변환을 **한 번에 하나씩** 지나가게 한다.
#:
#: 2026-09-03 job-28 에서 확정본 PDF 가 `Word COM: <unknown>.Open` 으로 실패해
#: 파일서버에 docx 만 올라갔다. 같은 초에 미리보기 변환이 돌고 있어 충돌을 의심했지만,
#: **동시 24건을 걸어도 재현되지 않았다**(2026-09-04). 원인을 확정하지 못했으므로
#: 잠금은 예방이고, 실제로 그 건을 살리는 것은 아래 재시도다.
_WORD_LOCK = threading.Lock()
#: 앞 변환이 끝나기를 기다리는 한계(초). 한 건이 보통 1~2초라 넉넉하다.
_WORD_WAIT = 180
#: Word 가 한 번 튕겨도 문서를 잃지 않도록 다시 해 본다. 실패가 드물고 원인이
#: 확정되지 않아, 사유를 가리지 않고 짧게 쉬었다 재시도한다.
_WORD_TRIES = 3


def docx_to_pdf_bytes(docx_bytes: bytes) -> tuple[bytes | None, str]:
    """docx 바이트를 PDF 바이트로 변환(Word COM). (bytes, error).

    docx2pdf 는 내부적으로 Word 를 COM 자동화하므로, 워커 스레드에서는 COM 초기화가 필요하다.
    **여러 스레드가 동시에 들어오면 안 된다** — _WORD_LOCK 참조.
    """
    tmpdir = tempfile.mkdtemp(prefix="rn_pdf_")
    in_path = os.path.join(tmpdir, "in.docx")
    out_path = os.path.join(tmpdir, "in.pdf")
    with open(in_path, "wb") as f:
        f.write(docx_bytes)

    if not _WORD_LOCK.acquire(timeout=_WORD_WAIT):
        return None, f"다른 PDF 변환이 {_WORD_WAIT}초 넘게 끝나지 않았습니다."
    co_inited = False
    try:
        try:
            import pythoncom  # pywin32

            pythoncom.CoInitialize()
            co_inited = True
        except Exception:  # noqa: BLE001
            pass

        from docx2pdf import convert

        last = ""
        for attempt in range(_WORD_TRIES):
            if attempt:
                time.sleep(1.5 * attempt)      # Word 가 정리될 틈을 준다
            try:
                if os.path.exists(out_path):
                    os.remove(out_path)
                convert(in_path, out_path)
            except Exception as e:  # noqa: BLE001
                last = f"PDF 변환 실패(Word COM): {e}"
                continue
            if os.path.exists(out_path):
                with open(out_path, "rb") as f:
                    return f.read(), ""
            last = "PDF 변환 결과 파일이 없습니다(Word 실행/권한 확인)."
        return None, last + f" ({_WORD_TRIES}번 시도)"
    finally:
        if co_inited:
            try:
                import pythoncom

                pythoncom.CoUninitialize()
            except Exception:  # noqa: BLE001
                pass
        _WORD_LOCK.release()


def render_rn_pdf(context: dict) -> tuple[bytes | None, str]:
    """Word 템플릿 채움 → PDF 변환까지. (pdf_bytes, error)."""
    docx_bytes, err = render_rn_docx(context)
    if docx_bytes is None:
        return None, err
    return docx_to_pdf_bytes(docx_bytes)


def fill_excel_template(
    rows: list[list], start_row: int = 2, sheet: str | None = None
) -> tuple[bytes | None, str]:
    """rn_excel 템플릿을 열어 start_row 부터 이슈 행을 채워 xlsx 바이트 반환. (bytes, error).

    (열 매핑은 템플릿 확정 시 정교화. 지금은 A열부터 순서대로 기록.)
    """
    path = T.path_for("rn_excel")
    if not path or not path.exists():
        return None, "Excel 템플릿(rn_excel)이 없습니다. [템플릿 관리]에서 업로드하세요."
    try:
        from openpyxl import load_workbook

        wb = load_workbook(str(path))
        ws = wb[sheet] if (sheet and sheet in wb.sheetnames) else wb.active
        for i, row in enumerate(rows):
            for j, val in enumerate(row):
                ws.cell(row=start_row + i, column=1 + j, value=val)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue(), ""
    except Exception as e:  # noqa: BLE001
        return None, f"Excel 템플릿 채우기 실패: {e}"
